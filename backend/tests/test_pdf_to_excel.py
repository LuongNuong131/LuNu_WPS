from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from app.document_ai.pipeline import PDFIntelligencePipeline
from app.processors.pdf_to_excel import PDFToExcelProcessor, _typed_excel_value


def make_native_pdf(path: Path) -> None:
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    table = Table([["Item", "Amount"], ["Service", "1.250.000 VNĐ"], ["Tax", "10%"]], colWidths=[70 * mm, 70 * mm])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    doc.build([Paragraph("Invoice INV-2026-001", styles["Heading1"]), Paragraph("Contact: test@example.com", styles["BodyText"]), Spacer(1, 8), table])


def test_typed_value_is_conservative() -> None:
    assert _typed_excel_value("1.250.000 VNĐ") == 1250000
    assert _typed_excel_value("1,250,000.50") == 1250000.5
    assert _typed_excel_value("INV-2026-001") == "INV-2026-001"


def test_native_pipeline_has_stable_evidence(tmp_path: Path) -> None:
    source = tmp_path / "invoice.pdf"
    make_native_pdf(source)
    document = PDFIntelligencePipeline().analyze(str(source), {"ocr_fallback": True, "ocr_language": "auto"})
    assert document.page_count == 1
    assert document.source and len(document.source.fingerprint or "") == 64
    assert document.pages[0].reading_order
    assert document.provenance
    assert document.state.value in {"ready", "review_required"}
    assert document.pipeline_version.pipeline_version == "5.0.0"
    assert document.quality_dimensions.evidence_coverage is not None
    assert document.facts
    assert all(fact.status.value in {"extracted", "normalized"} for fact in document.facts)
    assert json.dumps(document.diagnostics(), default=str)


def test_scan_pdf_uses_page_aware_ocr(tmp_path: Path) -> None:
    image = Image.new("RGB", (1600, 600), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 80), "Invoice INV-2026-009", fill="black")
    draw.text((80, 180), "Tong tien: 1.250.000 VNĐ", fill="black")
    draw.text((80, 280), "Email: test@example.com", fill="black")
    source = tmp_path / "scan.pdf"
    image.save(source, "PDF", resolution=200)
    document = PDFIntelligencePipeline().analyze(str(source), {"ocr_fallback": True, "ocr_language": "eng+vie"})
    assert document.pages[0].ocr_used is True
    assert document.pages[0].image_quality["selected_pass"] in {6, 11}
    assert "Invoice" in document.pages[0].raw_text or "INV" in document.pages[0].raw_text


def test_pdf_to_excel_contains_text_tables_and_audit(tmp_path: Path) -> None:
    source = tmp_path / "invoice.pdf"
    output = tmp_path / "result.xlsx"
    make_native_pdf(source)
    processor = PDFToExcelProcessor()
    assert processor.process([str(source)], str(output), {"extraction": "auto", "include_text": True, "ocr_fallback": True, "ocr_language": "auto", "use_canonical": True})
    workbook = load_workbook(output, data_only=False)
    assert "Overview" in workbook.sheetnames
    assert "Document Text" in workbook.sheetnames
    assert "Audit" in workbook.sheetnames
    assert any(name.startswith("Page") for name in workbook.sheetnames)
    assert workbook["Audit"].max_row > 1
    assert processor.last_diagnostics["audit_available"] is True
    assert processor.last_diagnostics["pages"] == 1
    assert processor.last_diagnostics["document_state"] in {"ready", "review_required"}
    assert "quality_dimensions" in processor.last_diagnostics
    assert "validation_results" in processor.last_diagnostics
    assert processor.last_diagnostics["facts"] >= 1
    assert "graph_edges" in processor.last_diagnostics
