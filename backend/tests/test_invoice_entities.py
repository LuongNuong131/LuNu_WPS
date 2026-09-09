from pathlib import Path

from openpyxl import load_workbook

from app.document_ai.models import BoundingBox, CanonicalDocument, Confidence, DocumentBlock, DocumentPage
from app.document_ai.semantics import extract_entities, parse_financial_number
from app.processors.pdf_to_excel import PDFToExcelProcessor


def test_localized_financial_parser_handles_vietnamese_and_international() -> None:
    assert parse_financial_number("1.000.000 VNĐ") == 1000000.0
    assert parse_financial_number("1,000,000.00 USD") == 1000000.0
    assert parse_financial_number("1.234,56 €") == 1234.56


def test_invoice_kvps_keep_block_provenance() -> None:
    block = DocumentBlock(
        "text", "Invoice No: INV-2026-001 | Date: 09/09/2026 | MST: 0312345678 | VAT: 10% | Subtotal: 1.000.000 VNĐ | Grand Total: 1.100.000 VNĐ",
        BoundingBox(10, 20, 500, 40), 1, 1, Confidence.from_score(.94, "test"), block_id="p1-b1",
    )
    document = CanonicalDocument("invoice.pdf", 1, pages=[DocumentPage(1, 600, 800, blocks=[block])])
    entities = extract_entities(document)
    by_type = {entity.entity_type: entity for entity in entities}
    assert by_type["Invoice_Number"].normalized_value == "INV-2026-001"
    assert by_type["Tax_Code"].normalized_value == "0312345678"
    assert by_type["Subtotal"].normalized_value == 1000000.0
    assert by_type["Grand_Total"].evidence[0].bbox.as_list() == [10, 20, 500, 40]


def test_workbook_contains_invoice_focused_sheets(tmp_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    source = tmp_path / "invoice.pdf"
    SimpleDocTemplate(str(source), pagesize=A4).build([Table([["Description", "Qty", "Total"], ["Widget", "2", "200"]], colWidths=[70 * mm, 30 * mm, 40 * mm], style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))])
    output = tmp_path / "invoice.xlsx"
    assert PDFToExcelProcessor().process([str(source)], str(output), {"extraction": "auto", "include_text": True, "ocr_fallback": True, "ocr_language": "auto", "use_canonical": True})
    workbook = load_workbook(output)
    assert "Overview" in workbook.sheetnames
    assert "Line Items" in workbook.sheetnames
    assert "Audit" in workbook.sheetnames
